"""Safe text extraction and chunking for uploaded documents.

Uploads are validated by extension + file signature before parsing. PDFs are
text-layer only; scanned images produce a clear failure instead of empty chunks.
Chunking prefers paragraph/sentence boundaries and keeps overlap for context.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}
EXTENSION_LABELS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
}
MAX_TEXT_LENGTH = 8_000_000


class DocumentParseError(ValueError):
    pass


def normalize_filename(filename: str) -> str:
    name = Path(filename or "未命名文件").name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "未命名文件"


def file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    if not data:
        raise DocumentParseError("文件内容为空")
    if len(data) > max_bytes:
        raise DocumentParseError(f"文件超过 {max_bytes // 1024 // 1024}MB 限制")
    ext = file_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentParseError("仅支持 PDF、DOCX、Markdown、TXT 文件")
    # Extensions alone are user-controlled. Check a small magic signature for
    # binary formats so renaming an executable to .pdf is rejected early.
    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentParseError("文件内容与 PDF 格式不符")
    elif ext == ".docx":
        if not data.startswith(b"PK\x03\x04") or not zipfile.is_zipfile(io.BytesIO(data)):
            raise DocumentParseError("文件内容与 DOCX 格式不符")
    else:
        if b"\x00" in data[:4096]:
            raise DocumentParseError("文件不是有效的文本文件")
    return ext


def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = file_extension(filename)
    if ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception as exc:
                    raise DocumentParseError("PDF 已加密，无法解析") from exc
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(pages)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("PDF 解析失败，请确认文件未损坏") from exc
    elif ext == ".docx":
        try:
            doc = DocxDocument(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)
        except Exception as exc:
            raise DocumentParseError("DOCX 解析失败，请确认文件未损坏") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise DocumentParseError("文本编码无法识别，请使用 UTF-8") from exc

    text = _clean_text(text)
    if not text:
        raise DocumentParseError("未能从文件中提取到文字（扫描版 PDF 暂不支持）")
    if len(text) > MAX_TEXT_LENGTH:
        raise DocumentParseError("文档文字量过大，请拆分后上传")
    return text


def _find_break(text: str, start: int, end: int) -> int:
    """Prefer a paragraph/sentence boundary so chunks stay readable."""
    floor = start + max(80, (end - start) // 2)
    for marker in ("\n\n", "\n", "。", "！", "？", ". ", "; ", "；"):
        idx = text.rfind(marker, floor, end)
        if idx != -1:
            return idx + len(marker)
    return end


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 80) -> list[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    overlap = max(0, min(overlap, chunk_size // 2))
    text = text.strip()
    if not text:
        return []

    chunks: list[str] = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        if end < length:
            end = _find_break(text, start, end)
        piece = text[start:end].strip()
        if piece:
            if len(piece) < 20 and chunks:
                chunks[-1] = f"{chunks[-1]}\n{piece}".strip()
            else:
                chunks.append(piece)
        if end >= length:
            break
        # Carry a tail of the previous chunk into the next one so a sentence
        # split at the boundary still has enough context for retrieval.
        next_start = max(end - overlap, start + 1)
        # Never let an overlap consume the whole next window.
        if next_start <= start:
            next_start = end
        start = next_start
    return chunks
